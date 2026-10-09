import base64
import io
import os

from fastapi.testclient import TestClient

from app import resume as r
from app.main import app, get_resume_ai

FIXTURES = os.path.join(os.path.dirname(__file__), "fixtures")


def sample() -> r.Resume:
    return r.Resume(
        name="Asha Rao",
        contact=["asha@example.com", "github.com/asha"],
        sections=[
            r.ResumeSection(
                title="Projects",
                kind="projects",
                items=[
                    r.ResumeItem(title="Chat app", bullets=["Built a chat app with WebSockets and Redis"]),
                    r.ResumeItem(title="Payments API", bullets=["Wrote a REST API in **Spring Boot** for payments"]),
                ],
            ),
            r.ResumeSection(title="Skills", kind="skills", lines=["**Languages:** Java, Python, SQL"]),
        ],
    )


def repos() -> list[r.RepoEvidence]:
    return [
        r.RepoEvidence(name="payments-api", url="https://github.com/asha/payments-api", languages=["Java"],
                       readme="Payments service. Uses Spring Boot with PostgreSQL and idempotent retries via Kafka."),
        r.RepoEvidence(name="raft-kv", url="https://github.com/asha/raft-kv", languages=["C++"],
                       readme="A Raft-replicated key-value store in C++17 with leader election and log compaction."),
    ]


def test_pdf_extraction_keeps_spaces_bold_and_links():
    with open(os.path.join(FIXTURES, "bold_links_resume.pdf"), "rb") as f:
        text, kind = r.extract_text("cv.pdf", f.read())
    assert kind == "text"
    assert "**Raft-replicated C++17** key-value cluster" in text  # spaces around bold survive
    assert "**Languages:** Java, C++, Python" in text
    assert "[GitHub](https://github.com/asha/raft-kv)" in text
    assert "• •" not in text and "**Asha Rao**" in text


def test_check_plan_keeps_real_edits_and_flags_invented_skills():
    plan = r.TailorPlan(
        summary="s",
        edits=[
            r.TailorEdit(section="Projects", before="Wrote a REST API in Spring Boot for payments",
                         after="Built a payments REST API in **Spring Boot**", reason="markup-insensitive match"),
            r.TailorEdit(section="Projects", before="Built a chat app with WebSockets and Redis",
                         after="Built a chat app with WebSockets, Redis and Kubernetes", reason="k8s in JD"),
            r.TailorEdit(section="Projects", before="Led a team of 10", after="Led a team of 12", reason="x"),
        ],
    )
    edits, _ = r.check_plan(sample(), plan)
    assert [e.before for e in edits] == ["Wrote a REST API in **Spring Boot** for payments",
                                         "Built a chat app with WebSockets and Redis"]
    assert edits[0].unsupported is False
    assert edits[1].unsupported is True and "Kubernetes" in edits[1].new_terms


def test_github_edits_must_quote_their_repository_and_arent_flagged_for_it():
    plan = r.TailorPlan(summary="s", edits=[
        # Real quote: Kafka comes from the repo, so it's supported.
        r.TailorEdit(section="Projects", before="Wrote a REST API in **Spring Boot** for payments",
                     after="Wrote a REST API in **Spring Boot** for payments with idempotent retries via Kafka",
                     reason="JD wants Kafka", source="github", repo="payments-api", evidence="idempotent retries via Kafka"),
        # Invented quote: dropped.
        r.TailorEdit(section="Projects", before="Built a chat app with WebSockets and Redis",
                     after="Built a chat app serving 1M users", reason="x", source="github", repo="payments-api",
                     evidence="serves 1M users"),
    ])
    edits, _ = r.check_plan(sample(), plan, "Kafka experience", repos())
    assert len(edits) == 1
    assert edits[0].source == "github" and edits[0].repo_url.endswith("payments-api")
    assert edits[0].unsupported is False


def test_swaps_come_from_a_real_repo_and_get_its_link():
    plan = r.TailorPlan(summary="s", swaps=[
        r.ProjectSwap(section="Projects", replace_title="Chat app", repo="raft-kv", reason="distributed systems role",
                      new_item=r.ResumeItem(title="Raft KV Store", subtitle="C++17",
                                            bullets=["Built a Raft-replicated key-value store with leader election"])),
        r.ProjectSwap(section="Projects", replace_title="Nonexistent", repo="raft-kv", reason="x",
                      new_item=r.ResumeItem(title="X", bullets=["y"])),
    ])
    _, swaps = r.check_plan(sample(), plan, "", repos())
    assert len(swaps) == 1 and swaps[0].id == 101
    assert swaps[0].new_item.links[0].url == "https://github.com/asha/raft-kv"
    assert swaps[0].unsupported is False


def test_link_additions_match_projects_to_repos():
    adds = r.link_additions(sample(), repos())
    assert [(a.item_title, a.url) for a in adds] == [("Payments API", "https://github.com/asha/payments-api")]


def test_apply_and_render_keeps_bold_links_and_escapes():
    resume = sample()
    resume.sections[1].lines = ["**Languages:** C++, C#, 100% & more_things"]
    applied = r.apply_to_resume(
        resume,
        [r.TailorEdit(section="Projects", before="Built a chat app with WebSockets and Redis",
                      after="Built a real-time chat app with WebSockets and Redis", reason="x")],
        [r.SectionOrder(section="projects", item_titles=["Payments API", "Chat app"])],
        links=r.link_additions(resume, repos()),
    )
    assert [i.title for i in applied.sections[0].items] == ["Payments API", "Chat app"]
    tex = r.render_latex(applied)
    assert r"\textbf{Languages:} C++, C\#, 100\% \& more\_things" in tex
    assert r"\href{https://github.com/asha/payments-api}" in tex
    assert "real-time chat app" in tex and tex.rstrip().endswith(r"\end{document}")
    assert "10pt" in r.render_latex(applied, compact=1) and "0.35in" in r.render_latex(applied, compact=2)


def test_ordinary_rewording_isnt_flagged_but_new_tools_are():
    text = "Flask service in Python with Redis and a Node.js backend"
    terms = r.introduced_terms("x", "Tested and fixed bugs in a Python-based back-end service", text)
    assert "Python-based" not in terms and "back-end" not in terms
    assert not any(r._looks_like_claim(t, first_word=(t == "Tested")) for t in terms)
    assert r._looks_like_claim("NoSQL") and r._looks_like_claim("Kubernetes") and r._looks_like_claim("CI/CD")
    assert not r._looks_like_claim("knowledge")


def test_wording_lifted_from_the_job_ad_is_flagged():
    plan = r.TailorPlan(summary="s", edits=[
        r.TailorEdit(section="Projects", before="Built a chat app with WebSockets and Redis",
                     after="Solved bottlenecks by building a chat app with WebSockets and Redis", reason="x"),
    ])
    edit = r.check_plan(sample(), plan, "You will solve performance bottlenecks.")[0][0]
    assert edit.unsupported is True and "bottlenecks" in edit.from_job_ad


def test_apply_to_own_tex_keeps_design_and_reports_misses():
    source = r"\begin{document}\item Wrote a REST API in \textbf{Spring Boot} for payments \& refunds\end{document}"
    edits = [
        r.TailorEdit(section="P", before="Wrote a REST API in **Spring Boot** for payments & refunds",
                     after="Built a payments & refunds REST API in **Spring Boot**", reason="x"),
        r.TailorEdit(section="P", before="Not in the file", after="whatever", reason="x"),
    ]
    out, missed = r.apply_to_tex(source, edits)
    assert r"Built a payments \& refunds REST API in \textbf{Spring Boot}" in out
    assert missed == [1]


def test_extract_reads_docx_bold_and_detects_tex():
    import docx

    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Languages: ").bold = True
    p.add_run("Java, Python. " * 10)
    buf = io.BytesIO()
    d.save(buf)
    text, kind = r.extract_text("cv.docx", buf.getvalue())
    assert kind == "text" and "**Languages:** Java" in text

    _, kind = r.extract_text("cv.txt", rb"\documentclass{article}\begin{document}hi\end{document}")
    assert kind == "tex"


def test_tailor_endpoint_checks_edits_and_honours_model_choice():
    seen = {}

    class FakeAi:
        model = "fake"

        async def structure(self, text):
            return sample(), None

        async def tailor(self, resume, jd, title, repo_list, instructions, one_page, model=None):
            seen.update(model=model, instructions=instructions, one_page=one_page, repos=[x.name for x in repo_list])
            return r.TailorPlan(
                summary="ok",
                edits=[r.TailorEdit(section="Skills", before="**Languages:** Java, Python, SQL",
                                    after="**Languages:** Java, SQL, Python", reason="x"),
                       r.TailorEdit(section="Skills", before="made up", after="also made up", reason="x")],
                matched_keywords=["Java"], missing_keywords=["Kafka"],
            ), None, model

    app.dependency_overrides[get_resume_ai] = lambda: FakeAi()
    try:
        client = TestClient(app)
        res = client.post("/resume/tailor", json={
            "resume": sample().model_dump(), "job_title": "Backend Intern",
            "job_description": "We want Java and Kafka experience for backend services at scale.",
            "repos": [x.model_dump() for x in repos()], "instructions": "Keep it to one page", "model": "flash",
        })
        assert res.status_code == 200, res.text
        body = res.json()
        assert len(body["edits"]) == 1 and body["dropped_edits"] == 1
        assert body["link_additions"][0]["repo"] == "payments-api"
        assert seen["model"] == "gemini-2.5-flash" and seen["instructions"] == "Keep it to one page"
        assert body["repos_used"] == ["payments-api", "raft-kv"]

        assert client.post("/resume/tailor", json={"resume": sample().model_dump(), "job_description": "x" * 50,
                                                   "model": "gpt"}).status_code == 422
        res = client.post("/resume/extract", json={"filename": "a.txt", "data_base64": base64.b64encode(b"hi").decode()})
        assert res.status_code == 422
    finally:
        app.dependency_overrides.clear()


def test_apply_kit_flags_claims_the_sources_never_make():
    from app import applykit

    repo = r.RepoEvidence(name="raft-kv", url="https://github.com/asha/raft-kv",
                          readme="A Raft-replicated key-value store in C++17 with leader election.")
    plan = applykit.KitPlan(answers=[
        applykit.KitAnswer(question="Why you?", answer="I built a Raft-replicated store in C++17 with leader election.",
                           grounded_on="raft-kv"),
        applykit.KitAnswer(question="Tell us more", answer="I also ran Kubernetes clusters at Infosys for 3 years.",
                           grounded_on="resume"),
        applykit.KitAnswer(question="Empty one", answer="   ", grounded_on=""),
    ])
    checked = applykit.check_answers(plan, sample(), [repo], "Available from June 2027.")

    assert len(checked) == 2  # the empty answer is dropped
    assert checked[0].unsupported is False
    assert checked[1].unsupported is True
    assert any(t in checked[1].new_terms for t in ("Kubernetes", "Infosys"))
