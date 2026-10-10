"""Resume tailoring: read a resume, propose grounded edits for one job, render LaTeX, compile a PDF.

The contract, same as everywhere else in RepoMatch: the model may only reword, reorder and emphasise what the resume
or the person's own GitHub repositories already say. Every edit names the exact text it replaces (checked against the
resume); every claim taken from GitHub quotes the repository text it rests on (checked against that text); and any
wording that brings in a tool, or copies the job ad, is flagged so the person decides, not the model.

Formatting survives the round trip through two lightweight markers inside strings: **bold** and [label](url).
"""

import asyncio
import io
import logging
import os
import re
import shutil
import tempfile
from typing import Protocol

from google import genai
from google.genai import types
from pydantic import BaseModel, Field

from app.schemas import TokenUsage

log = logging.getLogger("repomatch.ai")

MAX_RESUME_CHARS = 40_000
MAX_JD_CHARS = 20_000
MAX_README_CHARS = 2_500


# --- Structured resume -------------------------------------------------------


class Link(BaseModel):
    label: str
    url: str


class ResumeItem(BaseModel):
    title: str = ""
    subtitle: str = ""
    date: str = ""
    location: str = ""
    bullets: list[str] = []
    # e.g. [{"label": "GitHub", "url": "https://github.com/me/project"}]
    links: list[Link] = []


class ResumeSection(BaseModel):
    title: str
    # experience | projects | education | skills | achievements | other
    kind: str = "other"
    items: list[ResumeItem] = []
    # Free lines for sections without entries, e.g. "**Languages:** Java, Python".
    lines: list[str] = []


class Resume(BaseModel):
    name: str = ""
    headline: str = ""
    # Each entry may carry a link: "[github.com/me](https://github.com/me)".
    contact: list[str] = []
    sections: list[ResumeSection] = []


class RepoEvidence(BaseModel):
    """One of the person's repositories, as the job-relevant evidence the model may draw on."""

    name: str
    url: str
    description: str = ""
    languages: list[str] = []
    readme: str = ""

    def text(self) -> str:
        return "\n".join([self.name, self.description, " ".join(self.languages), self.readme])


_MARKUP_LINK = re.compile(r"\[([^\]]+)\]\((https?://[^)\s]+|mailto:[^)\s]+)\)")


def strip_markup(s: str) -> str:
    return _MARKUP_LINK.sub(r"\1", s or "").replace("**", "")


def _key(s: str) -> str:
    """How two pieces of resume text are compared: markup, spacing and case don't matter."""
    return re.sub(r"\s+", " ", strip_markup(s)).strip().lower()


def resume_strings(r: Resume) -> list[str]:
    """Every piece of text an edit is allowed to replace."""
    out: list[str] = []
    for s in r.sections:
        out.extend(s.lines)
        for it in s.items:
            out.extend(it.bullets)
    if r.headline:
        out.append(r.headline)
    return out


def resume_text(r: Resume) -> str:
    parts = [r.name, r.headline, " ".join(r.contact)]
    for s in r.sections:
        parts.append(s.title)
        parts.extend(s.lines)
        for it in s.items:
            parts.extend([it.title, it.subtitle, it.date, it.location, *it.bullets])
    return strip_markup("\n".join(p for p in parts if p))


# --- Extraction --------------------------------------------------------------

# A bullet glyph, or "-"/"*" followed by a space — never the "**" that opens a bold run.
_BULLET_GLYPHS = re.compile(r"^\s*(?:[•●▪◦‣∙·–—]\s*|[-*]\s+)(?=\S)")


def extract_text(filename: str, data: bytes) -> tuple[str, str]:
    """Returns (text, kind) where kind is "tex" when the upload is LaTeX source, else "text".

    PDFs keep what matters for a faithful re-render: word spacing, **bold** runs and [label](url) links."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        return _clean(_pdf_text(data)), "text"
    if name.endswith(".docx"):
        return _clean(_docx_text(data)), "text"
    text = data.decode("utf-8", errors="replace")
    kind = "tex" if name.endswith(".tex") or "\\documentclass" in text or "\\begin{document}" in text else "text"
    return (text if kind == "tex" else _clean(text))[:MAX_RESUME_CHARS * 2], kind


def _is_bold(fontname: str) -> bool:
    f = (fontname or "").lower()
    return any(k in f for k in ("bold", "black", "heavy", "semibold", "demi", "-bx", "cmbx"))


def _pdf_text(data: bytes) -> str:
    import pdfplumber

    out_lines: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            words = page.extract_words(x_tolerance=1, y_tolerance=3, keep_blank_chars=False, use_text_flow=True,
                                       extra_attrs=["fontname", "size"])
            links = [h for h in (page.hyperlinks or []) if h.get("uri")]
            # Group words into visual lines.
            lines: list[list[dict]] = []
            for w in words:
                if lines and abs(lines[-1][-1]["top"] - w["top"]) <= 3:
                    lines[-1].append(w)
                else:
                    lines.append([w])
            for line in lines:
                parts: list[str] = []
                bold_run: list[str] = []
                used_links: set[str] = set()

                def flush_bold():
                    if bold_run:
                        parts.append("**" + " ".join(bold_run) + "**")
                        bold_run.clear()

                for w in line:
                    link = next((h for h in links if h["x0"] - 1 <= w["x0"] and w["x1"] <= h["x1"] + 1
                                 and h["top"] - 2 <= w["top"] and w["bottom"] <= h["bottom"] + 2), None)
                    if link and link["uri"] not in used_links:
                        # The whole linked span becomes one [label](url).
                        label_words = [x["text"] for x in line if link["x0"] - 1 <= x["x0"] and x["x1"] <= link["x1"] + 1]
                        flush_bold()
                        parts.append(f"[{' '.join(label_words)}]({link['uri']})")
                        used_links.add(link["uri"])
                        continue
                    if link:
                        continue  # already emitted as part of the link label
                    if _is_bold(w.get("fontname", "")):
                        bold_run.append(w["text"])
                    else:
                        flush_bold()
                        parts.append(w["text"])
                flush_bold()
                out_lines.append(" ".join(parts))
    return "\n".join(_BULLET_GLYPHS.sub("- ", ln) if _BULLET_GLYPHS.match(ln) else ln for ln in out_lines)


def _docx_text(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    lines = []
    for p in document.paragraphs:
        chunks = []
        for run in p.runs:
            t = run.text
            if t.strip() and run.bold:
                chunks.append(f"**{t.strip()}**" + (" " if t.endswith(" ") else ""))
            else:
                chunks.append(t)
        lines.append("".join(chunks))
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace(" ", " ").replace("****", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_RESUME_CHARS]


# --- Gemini: structure and tailor ---------------------------------------------

STRUCTURE_INSTRUCTION = """\
You convert a resume into structured data, faithfully.

Rules:
- Copy text exactly as written. Do not rewrite, summarise, fix grammar or translate.
- Keep every bullet, in order. Keep every section, in order. Keep every achievement and skills line.
- Keep formatting markers exactly: **bold** and [label](url). For LaTeX input, turn \\textbf{x} into **x** and \
\\href{url}{label} into [label](url); drop other commands but keep their text.
- Don't include bullet characters (•, -) at the start of bullets.
- A link that belongs to an entry (e.g. "GitHub", "Live demo" next to a project title) goes in that item's links, \
not in its title.
- kind is one of: experience, projects, education, skills, achievements, other.
- Contact details (email, phone, LinkedIn, GitHub, website, city) go in contact, one per entry, keeping any link.
"""

TAILOR_INSTRUCTION = """\
You tailor a student's resume to one job. You are careful and honest: a recruiter will read this, and an \
interviewer will ask about every line.

Sources you may draw on: the resume itself, and — when provided — the person's own GitHub repositories.

You may:
- Reword an existing bullet or line so the relevant part leads and matches the job's vocabulary, where the resume \
already supports it.
- Strengthen a bullet with a concrete detail from the SAME project's repository (a tool, a component, a result) — \
only when the README or description states it. Then set source "github", repo to the repository name, and evidence \
to a short phrase copied word-for-word from that repository's text.
- Reorder items inside a section so the most relevant come first.
- Suggest swapping a weaker project for a repository that fits this job clearly better (swaps). The new item's title, \
subtitle (its tech) and 2-3 bullets must come only from that repository's text.

You must never:
- Add a skill, tool, metric, employer, responsibility or result that neither the resume nor the cited repository \
states.
- Change dates, titles, companies, schools or numbers. Touch contact details.
- Remove or alter formatting markers: keep **bold** spans and [label](url) links exactly as they are in "before".

Write like a strong engineer describing their own work, not like a keyword list:
- Plain, specific, active wording, one idea per bullet, not much longer than the original.
- Never bolt job-ad jargon onto a bullet ("a back-end framework (Spring Boot)"). Name the real tool.
- If a bullet is already clear and relevant, leave it alone. Fewer, better edits beat many small ones.

Follow the person's own instructions when given. If they ask for one page, prefer tightening wording to adding it, \
and don't make bullets longer.

For each edit, "before" must be copied character-for-character from one bullet or line of the resume. At most 8 \
edits and 2 swaps. List requirements the resume or repos already show in matched_keywords, and requirements neither \
shows in missing_keywords — gaps for the person, never text to insert.
"""


class TailorEdit(BaseModel):
    section: str
    before: str
    after: str
    reason: str
    # "resume" when the edit only rewords; "github" when it uses a detail from a repository.
    source: str = "resume"
    repo: str = ""
    evidence: str = ""


class SectionOrder(BaseModel):
    section: str
    item_titles: list[str]


class ProjectSwap(BaseModel):
    section: str
    replace_title: str
    repo: str
    new_item: ResumeItem
    reason: str


class TailorPlan(BaseModel):
    summary: str
    edits: list[TailorEdit] = Field(default_factory=list)
    reorder: list[SectionOrder] = Field(default_factory=list)
    swaps: list[ProjectSwap] = Field(default_factory=list)
    matched_keywords: list[str] = Field(default_factory=list)
    missing_keywords: list[str] = Field(default_factory=list)


class CheckedEdit(TailorEdit):
    id: int
    # Terms the new wording introduces that appear nowhere in the resume (or the cited repository).
    new_terms: list[str] = []
    # True when new_terms look like claims: a skill or tool name, or wording lifted from the job ad.
    unsupported: bool = False
    # The subset of new_terms copied from the job description rather than from the resume or repository.
    from_job_ad: list[str] = []
    repo_url: str = ""


class CheckedSwap(ProjectSwap):
    id: int
    repo_url: str
    new_terms: list[str] = []
    unsupported: bool = False


class BulletAddition(BaseModel):
    """A line the person adds themselves, to answer something the job asks for that the resume didn't show.

    The text is theirs. Nothing here is drafted from the job ad — the gap it answers is recorded only so the person
    can see why the line is there."""

    section: str
    # Empty to add the line to the section itself rather than to one of its items.
    item_title: str = ""
    text: str
    gap: str = ""


class LinkAddition(BaseModel):
    id: int
    section: str
    item_title: str
    repo: str
    url: str


_WORD = re.compile(r"[A-Za-z][A-Za-z0-9+#./-]*")
_COMMON = set(
    """a an and or the of to in on for with by from as at into via using used use built build developed develop
    designed design implemented implement created create led lead improved improve reduced reduce increased
    increase optimized optimize managed manage worked work across over under within including include
    which that this these those their its his her our your my we i it is are was were be been being has have had
    team teams project projects system systems service services application applications app apps feature features
    data user users performance scalable scale high low new end end-to-end real real-time time based driven
    efficient efficiency reliable reliability production across multiple several key core main""".split()
)


def _norm(word: str) -> str:
    """"Back-end" and "backend", "APIs" and "API" are the same word for this check."""
    w = word.lower().strip(".,;:()*").replace("-", "")
    return w[:-1] if len(w) > 3 and w.endswith("s") else w


def _vocabulary(*texts: str) -> set[str]:
    known: set[str] = set()
    for t in texts:
        for w in _WORD.findall(strip_markup(t)):
            known.add(_norm(w))
            known.update(_norm(part) for part in w.split("-") if part)
    return known


def introduced_terms(before: str, after: str, resume_all_text: str, extra_text: str = "") -> list[str]:
    """Words in `after` that appear neither in `before`, anywhere else on the resume, nor in `extra_text`."""
    known = _vocabulary(resume_all_text, before, extra_text)
    out: list[str] = []
    for w in _WORD.findall(strip_markup(after)):
        parts = [w] + [p for p in w.split("-") if p]
        if all(_norm(p) in known or _norm(p) in _COMMON or len(_norm(p)) < 2 for p in parts[1:] or parts):
            continue
        if _norm(w) in known or _norm(w) in _COMMON or len(_norm(w)) < 2:
            continue
        if _norm(w) not in (_norm(x) for x in out):
            out.append(w.strip(".,;:()*"))
    return out


def _looks_like_claim(term: str, first_word: bool = False) -> bool:
    """Tools and skills look like names: acronyms (AWS), inner capitals (NoSQL), symbols (C++, Node.js, CI/CD),
    digits (S3), or a capital letter mid-sentence. A capitalised first word ("Tested") is just a sentence start."""
    if any(c in term for c in "+#/") or any(c.isdigit() for c in term) or re.search(r"\w\.\w", term):
        return True
    if len(term) > 1 and (term.isupper() or any(c.isupper() for c in term[1:])):
        return True
    return term[:1].isupper() and not first_word


def _name_key(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", strip_markup(s).lower())


def _same_project(item_title: str, repo_name: str) -> bool:
    a, b = _name_key(item_title), _name_key(repo_name)
    if len(a) < 3 or len(b) < 3:
        return False
    return a == b or a.startswith(b) or b.startswith(a) or (len(b) >= 5 and b in a)


def _has_repo_link(item: ResumeItem) -> bool:
    blob = " ".join([item.title, item.subtitle] + [link.url for link in item.links])
    return "github.com" in blob.lower()


def link_additions(resume: Resume, repos: list[RepoEvidence], start_id: int = 201) -> list[LinkAddition]:
    """Projects on the resume that are clearly one of the person's repositories but don't link to it."""
    out: list[LinkAddition] = []
    for s in resume.sections:
        for it in s.items:
            if _has_repo_link(it):
                continue
            repo = next((r for r in repos if _same_project(it.title, r.name)), None)
            if repo:
                out.append(LinkAddition(id=start_id + len(out), section=s.title, item_title=strip_markup(it.title),
                                        repo=repo.name, url=repo.url))
    return out


def check_plan(resume: Resume, plan: TailorPlan, job_description: str = "", repos: list[RepoEvidence] | None = None):
    """Keeps edits whose `before` really is resume text; GitHub edits must quote their repository; flags wording that
    adds unstated claims or echoes the job ad. Returns (edits, swaps)."""
    repos = repos or []
    by_repo = {_name_key(r.name): r for r in repos}
    allowed = {_key(s): s for s in resume_strings(resume)}
    all_text = resume_text(resume)
    jd_words = {_norm(w) for w in _WORD.findall(job_description or "")}

    edits: list[CheckedEdit] = []
    seen: set[str] = set()
    for e in plan.edits:
        key = _key(e.before)
        if key not in allowed or key in seen or not e.after.strip() or _key(e.after) == key:
            continue
        repo = by_repo.get(_name_key(e.repo)) if e.source == "github" else None
        if e.source == "github":
            quote = _key(e.evidence)
            if not repo or len(quote) < 4 or quote not in _key(repo.text()):
                continue  # a repository claim without a real quote from that repository
        seen.add(key)
        extra = repo.text() if repo else ""
        terms = introduced_terms(e.before, e.after, all_text, extra)
        first = (_WORD.findall(strip_markup(e.after)) or [""])[0].strip(".,;:()*")
        echoed = [t for t in terms if len(_norm(t)) >= 5 and _norm(t) in jd_words]
        edits.append(CheckedEdit(
            id=len(edits) + 1, section=e.section, before=allowed[key], after=e.after.strip(), reason=e.reason,
            source="github" if repo else "resume", repo=repo.name if repo else "", evidence=e.evidence if repo else "",
            repo_url=repo.url if repo else "", new_terms=terms,
            unsupported=bool(echoed) or any(_looks_like_claim(t, first_word=(t == first)) for t in terms),
            from_job_ad=echoed,
        ))
        if len(edits) >= 8:
            break

    titles = {_key(it.title): it for s in resume.sections for it in s.items}
    swaps: list[CheckedSwap] = []
    for sw in plan.swaps:
        repo = by_repo.get(_name_key(sw.repo))
        if not repo or _key(sw.replace_title) not in titles or not sw.new_item.bullets:
            continue
        new_text = "\n".join([sw.new_item.title, sw.new_item.subtitle, *sw.new_item.bullets])
        terms = introduced_terms("", new_text, repo.text())
        item = sw.new_item.model_copy(deep=True)
        if not any("github.com" in link.url for link in item.links):
            item.links.append(Link(label="GitHub", url=repo.url))
        swaps.append(CheckedSwap(
            id=101 + len(swaps), section=sw.section, replace_title=titles[_key(sw.replace_title)].title, repo=repo.name,
            new_item=item, reason=sw.reason, repo_url=repo.url, new_terms=terms,
            unsupported=any(_looks_like_claim(t) for t in terms),
        ))
        if len(swaps) >= 2:
            break
    return edits, swaps


def _usage(resp) -> TokenUsage | None:
    if not resp.usage_metadata:
        return None
    return TokenUsage(
        prompt_tokens=resp.usage_metadata.prompt_token_count or 0,
        output_tokens=resp.usage_metadata.candidates_token_count or 0,
        thinking_tokens=resp.usage_metadata.thoughts_token_count or 0,
    )


PHRASE_INSTRUCTION = """You rewrite one sentence a person wrote about their own work into a resume bullet.

You are not a ghostwriter and not a salesperson. The facts are theirs and the facts are fixed:

- Use ONLY what their sentence says. Never add a tool, a number, a company, a scale, a result or a quality they did not write. If their sentence is vague, the bullet stays vague.
- Keep every specific they give: project names, counts, tools, timeframes. Dropping a detail they wrote is as wrong as adding one they didn't.
- Start with a plain past-tense verb (Built, Led, Wrote, Fixed, Ran, Taught). No "spearheaded", no "leveraged".
- One sentence, at most 28 words, no full stop at the end if their own lines have none.
- Never mention the job, the company applying to, or the quality being demonstrated. A bullet that says "showing strong collaboration" is worthless; the work itself shows it.
- If their sentence is already a good bullet, return it unchanged."""


class Phrasing(BaseModel):
    bullet: str = ""


class CheckedPhrasing(Phrasing):
    # Words the bullet uses that the person did not write and that appear nowhere on their resume.
    new_terms: list[str] = []
    # True when those words look like claims. The caller shows the person's own sentence instead.
    unsupported: bool = False


class ResumeAi(Protocol):
    model: str

    async def structure(self, text: str) -> tuple[Resume, TokenUsage | None]: ...

    async def phrase(self, sentence: str, resume: Resume,
                     model: str | None = None) -> tuple[Phrasing, TokenUsage | None, str]: ...

    async def tailor(self, resume: Resume, job_description: str, job_title: str, repos: list[RepoEvidence],
                     instructions: str, one_page: bool, model: str | None = None) -> tuple[TailorPlan, TokenUsage | None, str]: ...


def tailor_prompt(resume: Resume, job_description: str, job_title: str, repos: list[RepoEvidence], instructions: str,
                  one_page: bool) -> str:
    parts = [f"# Job: {job_title}\n\n{job_description[:MAX_JD_CHARS]}"]
    wishes = []
    if one_page:
        wishes.append("The final resume must fit on ONE page.")
    if instructions.strip():
        wishes.append(instructions.strip()[:1500])
    if wishes:
        parts.append("# The person's instructions\n\n" + "\n".join(f"- {w}" for w in wishes))
    parts.append(f"# Resume (structured)\n\n{resume.model_dump_json(indent=1)}")
    if repos:
        blocks = [
            f"## {r.name}\nurl: {r.url}\ndescription: {r.description}\nlanguages: {', '.join(r.languages)}\n"
            f"README:\n{r.readme[:MAX_README_CHARS]}"
            for r in repos
        ]
        parts.append("# The person's GitHub repositories (most relevant to this job first)\n\n" + "\n\n".join(blocks))
    return "\n\n".join(parts)


class GeminiResumeAi:
    def __init__(self, client: genai.Client, model: str):
        self._client = client
        self.model = model

    async def structure(self, text: str) -> tuple[Resume, TokenUsage | None]:
        resp = await self._client.aio.models.generate_content(
            model=self.model,
            contents=f"# Resume\n\n{text[:MAX_RESUME_CHARS]}",
            config=types.GenerateContentConfig(
                system_instruction=STRUCTURE_INSTRUCTION,
                temperature=0.0,
                response_mime_type="application/json",
                response_schema=Resume,
            ),
        )
        resume = resp.parsed if isinstance(resp.parsed, Resume) else Resume.model_validate_json(resp.text)
        return resume, _usage(resp)

    async def phrase(self, sentence, resume, model=None):
        used = model or self.model
        resp = await self._client.aio.models.generate_content(
            model=used,
            contents="Their sentence:\n" + sentence[:1500],
            config=types.GenerateContentConfig(
                system_instruction=PHRASE_INSTRUCTION,
                temperature=0.1,
                response_mime_type="application/json",
                response_schema=Phrasing,
            ),
        )
        parsed = resp.parsed if isinstance(resp.parsed, Phrasing) else Phrasing.model_validate_json(resp.text or "{}")
        return parsed, _usage(resp), used

    async def tailor(self, resume, job_description, job_title, repos, instructions, one_page, model=None):
        used = model or self.model
        resp = await self._client.aio.models.generate_content(
            model=used,
            contents=tailor_prompt(resume, job_description, job_title, repos, instructions, one_page),
            config=types.GenerateContentConfig(
                system_instruction=TAILOR_INSTRUCTION,
                temperature=0.2,
                response_mime_type="application/json",
                response_schema=TailorPlan,
            ),
        )
        plan = resp.parsed if isinstance(resp.parsed, TailorPlan) else TailorPlan.model_validate_json(resp.text)
        return plan, _usage(resp), used


# --- Applying changes --------------------------------------------------------


def apply_to_resume(resume: Resume, edits: list[TailorEdit], reorder: list[SectionOrder],
                    swaps: list[ProjectSwap] | None = None, links: list[LinkAddition] | None = None,
                    additions: list[BulletAddition] | None = None) -> Resume:
    r = resume.model_copy(deep=True)
    replace = {_key(e.before): e.after for e in edits}

    def swap_text(s: str) -> str:
        return replace.get(_key(s), s)

    for section in r.sections:
        section.lines = [swap_text(x) for x in section.lines]
        for it in section.items:
            it.bullets = [swap_text(b) for b in it.bullets]
        for sw in swaps or []:
            for i, it in enumerate(section.items):
                if _key(it.title) == _key(sw.replace_title):
                    section.items[i] = sw.new_item.model_copy(deep=True)
        for add in links or []:
            for it in section.items:
                if _key(it.title) == _key(add.item_title) and not _has_repo_link(it):
                    it.links.append(Link(label="GitHub", url=add.url))
        order = next((o for o in reorder if o.section.strip().lower() == section.title.strip().lower()), None)
        if order:
            rank = {_key(t): i for i, t in enumerate(order.item_titles)}
            section.items.sort(key=lambda it: rank.get(_key(it.title), len(rank)))
    for add in additions or []:
        _add_bullet(r, add)
    if r.headline:
        r.headline = swap_text(r.headline)
    return r


def _add_bullet(r: Resume, add: BulletAddition) -> None:
    """Puts the person's own line where they asked for it, creating the section if it isn't there yet."""
    text = (add.text or "").strip()
    if not text:
        return
    section = next((s for s in r.sections if _key(s.title) == _key(add.section)), None)
    if section is None:
        section = ResumeSection(title=add.section.strip() or "Additional", kind="other", items=[], lines=[])
        r.sections.append(section)
    if add.item_title:
        item = next((it for it in section.items if _key(it.title) == _key(add.item_title)), None)
        if item is not None:
            item.bullets.append(text)
            return
    # No entry named, or the named one is gone (swapped out, say). The line goes to the section itself: a line about
    # one project filed under another would make the resume say something false about that other project.
    section.lines.append(text)


def check_phrasing(sentence: str, phrasing: Phrasing, resume: Resume) -> CheckedPhrasing:
    """A tightened bullet may only use the person's own words plus what's already on their resume."""
    bullet = (phrasing.bullet or "").strip()
    if not bullet:
        return CheckedPhrasing(bullet=sentence.strip())
    new = introduced_terms(sentence, bullet, resume_text(resume))
    claims = [t for i, t in enumerate(new) if _looks_like_claim(t, first_word=(i == 0 and bullet.startswith(t)))]
    if claims:
        # It added something they didn't say: hand back their own sentence and say why.
        return CheckedPhrasing(bullet=sentence.strip(), new_terms=claims, unsupported=True)
    return CheckedPhrasing(bullet=bullet, new_terms=new)


def latex_variants(text: str) -> list[str]:
    """How a resume string may appear in LaTeX source."""
    plain = strip_markup(text)
    return list(dict.fromkeys([tex_inline(text), text, plain, latex_escape(plain),
                               latex_escape(plain).replace("\\textasciitilde{}", "\\~{}")]))


def apply_to_tex(source: str, edits: list[TailorEdit]) -> tuple[str, list[int]]:
    """Replaces each edit's text inside the person's own .tex. Returns the new source and indexes not applied."""
    out = source
    missed: list[int] = []
    for i, e in enumerate(edits):
        for variant in latex_variants(e.before):
            if variant and variant in out:
                out = out.replace(variant, tex_inline(e.after), 1)
                break
        else:
            missed.append(i)
    return out, missed


# --- LaTeX rendering ----------------------------------------------------------

_LATEX_SPECIALS = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}


def latex_escape(s: str) -> str:
    s = "".join(_LATEX_SPECIALS.get(c, c) for c in s)
    return (s.replace("–", "--").replace("—", "---").replace("’", "'").replace("“", "``")
            .replace("”", "''").replace("₹", "Rs. "))


_INLINE = re.compile(r"(\*\*.+?\*\*|\[[^\]]+\]\((?:https?://|mailto:)[^)\s]+\))")


def tex_inline(s: str) -> str:
    """Escapes text while turning **bold** into \\textbf and [label](url) into \\href."""
    out: list[str] = []
    for part in _INLINE.split(s or ""):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append(r"\textbf{" + tex_inline(part[2:-2]) + "}")
        elif (m := _MARKUP_LINK.fullmatch(part)):
            url = m.group(2).replace("%", r"\%").replace("#", r"\#")
            out.append(r"\href{" + url + "}{" + latex_escape(m.group(1)) + "}")
        else:
            out.append(latex_escape(part))
    return "".join(out)


def preamble(compact: int = 0) -> str:
    """compact 0 = normal; 1 and 2 tighten font size, margins and spacing to fit one page."""
    size, margin, item_sep, section_before, entry_gap = [
        ("11pt", "0.55in", "0.5pt", "9pt", "4pt"),
        ("10pt", "0.45in", "0pt", "6pt", "2pt"),
        ("10pt", "0.35in", "0pt", "4pt", "1pt"),
    ][min(compact, 2)]
    tighter = r"\linespread{0.95}" if compact >= 2 else ""
    return rf"""\documentclass[{size},a4paper]{{article}}
\usepackage[margin={margin}]{{geometry}}
\usepackage{{fontspec}}
% No fi/fl ligatures: they are one glyph in the file, and a parser that doesn't know them reads "prole".
\defaultfontfeatures{{Ligatures=NoCommon}}
\usepackage{{enumitem}}
\usepackage{{titlesec}}
\usepackage[hidelinks]{{hyperref}}
\usepackage{{xcolor}}
{tighter}
\pagestyle{{empty}}
\setlength{{\parindent}}{{0pt}}
\titleformat{{\section}}{{\large\bfseries\scshape}}{{}}{{0em}}{{}}[\vspace{{-5pt}}\rule{{\linewidth}}{{0.5pt}}]
\titlespacing*{{\section}}{{0pt}}{{{section_before}}}{{4pt}}
\setlist[itemize]{{leftmargin=1.2em, itemsep={item_sep}, topsep=1pt, parsep=0pt, partopsep=0pt}}
\newcommand{{\entry}}[2]{{\par\noindent\textbf{{#1}}\hfill #2\par}}
\newcommand{{\entrysub}}[2]{{\noindent\textit{{#1}}\hfill\textit{{#2}}\par}}
\newcommand{{\entrygap}}{{\vspace{{{entry_gap}}}}}
"""


def _contact(c: str) -> str:
    c = c.strip()
    if _MARKUP_LINK.search(c):
        return tex_inline(c)
    if re.fullmatch(r"[^@\s]+@[^@\s]+\.[a-z]{2,}", c, re.I):
        return rf"\href{{mailto:{c}}}{{{latex_escape(c)}}}"
    if re.match(r"(https?://)?(www\.)?(github|linkedin|gitlab|leetcode|codeforces)\.com/", c, re.I) or c.startswith("http"):
        url = c if c.startswith("http") else "https://" + c
        return rf"\href{{{url}}}{{{latex_escape(c.replace('https://', '').replace('http://', ''))}}}"
    return tex_inline(c)


def render_latex(r: Resume, compact: int = 0) -> str:
    body: list[str] = [preamble(compact), r"\begin{document}", r"\begin{center}"]
    body.append(rf"{{\LARGE\bfseries {latex_escape(strip_markup(r.name) or 'Your Name')}}}\\[3pt]")
    if r.headline:
        body.append(rf"{tex_inline(r.headline)}\\[2pt]")
    if r.contact:
        body.append(r" \textbar{} ".join(_contact(c) for c in r.contact))
    body.append(r"\end{center}")

    for s in r.sections:
        body.append(rf"\section*{{{latex_escape(strip_markup(s.title))}}}")
        for i, it in enumerate(s.items):
            if i > 0:
                body.append(r"\entrygap")
            title = tex_inline(it.title)
            if it.links:
                title += r" \textbar{} " + r" \textbar{} ".join(
                    rf"\href{{{link.url}}}{{\normalfont\underline{{{latex_escape(link.label)}}}}}" for link in it.links)
            if it.title or it.date:
                body.append(rf"\entry{{{title}}}{{{tex_inline(it.date)}}}")
            if it.subtitle or it.location:
                body.append(rf"\entrysub{{{tex_inline(it.subtitle)}}}{{{tex_inline(it.location)}}}")
            if it.bullets:
                body.append(r"\begin{itemize}")
                body.extend(rf"  \item {tex_inline(_BULLET_GLYPHS.sub('', b))}" for b in it.bullets)
                body.append(r"\end{itemize}")
        if s.lines:
            if s.kind == "achievements":
                body.append(r"\begin{itemize}")
                body.extend(rf"  \item {tex_inline(_BULLET_GLYPHS.sub('', line))}" for line in s.lines)
                body.append(r"\end{itemize}")
            else:
                body.append(r"\\ ".join(tex_inline(_BULLET_GLYPHS.sub("", line)) for line in s.lines) + r"\par")
    body.append(r"\end{document}")
    return "\n".join(body) + "\n"


# --- Compiling ---------------------------------------------------------------


class CompileResult(BaseModel):
    ok: bool
    pdf: bytes | None = None
    log: str = ""
    pages: int = 0


def page_count(pdf: bytes) -> int:
    from pypdf import PdfReader

    return len(PdfReader(io.BytesIO(pdf)).pages)


async def compile_pdf(latex: str, timeout_s: int = 150) -> CompileResult:
    """Compiles with Tectonic (self-contained TeX engine; package bundle baked into the image)."""
    tectonic = shutil.which("tectonic")
    if not tectonic:
        return CompileResult(ok=False, log="Tectonic isn't installed on this server.")
    with tempfile.TemporaryDirectory() as work:
        src = os.path.join(work, "resume.tex")
        with open(src, "w", encoding="utf-8") as f:
            f.write(latex)
        proc = await asyncio.create_subprocess_exec(
            tectonic, "--chatter", "minimal", "--keep-logs", "--outdir", work, src,
            stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.STDOUT,
        )
        try:
            out, _ = await asyncio.wait_for(proc.communicate(), timeout=timeout_s)
        except TimeoutError:
            proc.kill()
            return CompileResult(ok=False, log=f"LaTeX compile timed out after {timeout_s}s")
        log_text = out.decode("utf-8", errors="replace")[-6000:]
        pdf_path = os.path.join(work, "resume.pdf")
        if proc.returncode != 0 or not os.path.exists(pdf_path):
            return CompileResult(ok=False, log=log_text)
        with open(pdf_path, "rb") as f:
            pdf = f.read()
        return CompileResult(ok=True, pdf=pdf, log=log_text, pages=page_count(pdf))
